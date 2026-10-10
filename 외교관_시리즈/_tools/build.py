#!/usr/bin/env python3
"""원고(.txt) → 대본(.docx) 변환 및 검증.

사용법:
  python3 build.py check  <원고.txt> [...]   # 형식·턴 수 검증만
  python3 build.py build  <원고.txt> [...]   # 검증 후 docx 생성
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parent.parent
TURNS = 150
META_KEYS = ["part", "ep", "title", "title_ko", "tagline", "info", "cast", "note", "history", "preview"]
REQUIRED = ["part", "ep", "title", "title_ko", "tagline", "info", "cast", "note", "history", "preview"]

PALETTE = ["0B5394", "AA3939", "38761D", "7F6000", "2E75B6", "741B47", "B45F06",
           "134F5C", "674EA7", "990000", "45818E", "BF9000", "6AA84F", "A64D79"]
FIXED = {"FALK": "1F3864", "JAN": "38761D", "WINTER": "741B47", "MIRA": "134F5C",
         "ANNA": "B45F06", "SEDOV": "990000", "ELISABETH": "A64D79", "GEORG": "7F6000",
         "CLAIRE": "674EA7", "NOWAK": "0B5394", "HALE": "45818E", "GREVE": "BF9000",
         "BRÜCKNER": "AA3939", "LEHNERT": "2E75B6"}


def parse(path):
    meta = {"cast": [], "note": [], "history": []}
    sections, errors = [], []
    pending = None  # (speaker, german) waiting for Korean line
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    for no, raw in enumerate(lines, 1):
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("@"):
            key, _, val = line[1:].partition(" ")
            if key not in META_KEYS:
                errors.append(f"{no}: 알 수 없는 메타 키 @{key}")
            elif key in ("cast", "note", "history"):
                meta[key].append(val.strip())
            else:
                meta[key] = val.strip()
            continue
        if line.startswith("== "):
            if pending:
                errors.append(f"{no}: 앞 대사에 한국어 번역(>) 누락")
                pending = None
            sections.append([line[3:].strip(), []])
            continue
        if line.startswith(">"):
            if not pending:
                errors.append(f"{no}: 대사 없는 번역 줄")
                continue
            sections[-1][1].append((pending[0], pending[1], line[1:].strip()))
            pending = None
            continue
        m = re.match(r"^([A-ZÄÖÜ][A-ZÄÖÜ\-]*):\s+(.+)$", line)
        if m:
            if pending:
                errors.append(f"{no}: 앞 대사에 한국어 번역(>) 누락")
            if not sections:
                errors.append(f"{no}: 장면(== ) 전에 대사가 있음")
                sections.append(["?", []])
            pending = (m.group(1), m.group(2).strip())
            continue
        errors.append(f"{no}: 해석할 수 없는 줄: {line[:60]}")
    if pending:
        errors.append("끝: 마지막 대사에 번역 누락")
    for k in REQUIRED:
        if not meta.get(k):
            errors.append(f"메타 @{k} 누락")
    cast_names = {c.split("|")[0].strip() for c in meta["cast"]}
    turns = [t for _, ts in sections for t in ts]
    if len(turns) != TURNS:
        errors.append(f"턴 수 {len(turns)} ≠ {TURNS}")
    for sp in sorted({t[0] for t in turns} - cast_names):
        errors.append(f"등장인물 목록(@cast)에 없는 화자: {sp}")
    for i, t in enumerate(turns):
        if re.search(r"[가-힣]", t[1]):
            errors.append(f"턴 {i+1}: 독일어 대사에 한글 포함")
        if not re.search(r"[가-힣]", t[2]):
            errors.append(f"턴 {i+1}: 번역에 한글 없음")
    return meta, sections, errors


def para(doc, text="", *, bold=False, color=None, size=None, before=None, after=None, center=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if before is not None:
        pf.space_before = Twips(before)
    if after is not None:
        pf.space_after = Twips(after)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        run(p, text, bold=bold, color=color, size=size)
    return p


def run(p, text, *, bold=False, color=None, size=None):
    r = p.add_run(text)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if size:
        r.font.size = Pt(size)
    return r


def build(meta, sections, out):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)
    sec.top_margin = sec.bottom_margin = Twips(1134)
    sec.left_margin = sec.right_margin = Twips(1417)

    para(doc, "DER UNTERHÄNDLER", bold=True, color="808080", size=15, before=2400, after=200, center=True)
    para(doc, "독일어 낭독 교재 · 외교 스릴러", color="808080", size=13, before=0, after=100, center=True)
    para(doc, f"{meta['part']}부 · {meta['ep']}화", bold=True, color="1F3864", size=12, before=600, after=100, center=True)
    para(doc, meta["title"], bold=True, color="1F3864", size=20, before=0, after=40, center=True)
    para(doc, meta["tagline"], color="404040", size=12, before=600, after=40, center=True)
    para(doc, meta["info"], color="595959", size=9, before=0, after=40, center=True)
    para(doc, "지문 없이 대사만으로 구성 · 독일어 본문 + 한국어 대역", color="595959", size=9, before=0, after=100, center=True)
    para(doc, "등장인물", bold=True, color="1F3864", size=10, before=500, after=150, center=True)
    colors, k = {}, 0
    for c in meta["cast"]:
        name, _, desc = c.partition("|")
        name = name.strip()
        if name in FIXED:
            colors[name] = FIXED[name]
        else:
            colors[name] = PALETTE[k % len(PALETTE)]
            k += 1
        p = para(doc, after=40, center=True)
        run(p, name + "  ", bold=True, size=10)
        run(p, desc.strip(), color="595959", size=9)
    for n in meta["note"]:
        para(doc, n, color="808080", size=8, before=120, after=0)
    para(doc)
    for title, turns in sections:
        para(doc, title, bold=True, color="1F3864", size=12, before=400, after=240)
        for sp, de, ko in turns:
            p = para(doc, before=160, after=20)
            run(p, sp, bold=True, color=colors.get(sp, "1A1A1A"), size=10)
            run(p, "   " + de, color="1A1A1A", size=11)
            para(doc, ko, color="808080", size=9, after=0)
    para(doc, "역사 노트", bold=True, color="1F3864", size=12, before=500, after=100)
    for h in meta["history"]:
        para(doc, h, color="595959", size=9, after=80)
    para(doc, "다음 화 예고", bold=True, color="1F3864", size=12, before=500, after=100)
    para(doc, meta["preview"], color="595959", size=9)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def main():
    mode, files = sys.argv[1], sys.argv[2:]
    bad = 0
    for f in files:
        meta, sections, errors = parse(f)
        if errors:
            bad += 1
            print(f"✗ {f}")
            for e in errors[:40]:
                print("   ", e)
            continue
        n = sum(len(t) for _, t in sections)
        if mode == "build":
            safe = re.sub(r"[^A-Za-z0-9]+", "_", meta["title"].translate(
                str.maketrans({"ä": "ae", "ö": "oe", "ü": "ue", "Ä": "Ae", "Ö": "Oe", "Ü": "Ue", "ß": "ss"}))).strip("_")
            out = ROOT / f"{meta['part']}부" / f"{meta['part']}부_{meta['ep']}화_{safe}.docx"
            build(meta, sections, out)
            print(f"✓ {f} → {out.relative_to(ROOT)} ({n}턴)")
        else:
            print(f"✓ {f} ({n}턴, 장면 {len(sections)}개)")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
